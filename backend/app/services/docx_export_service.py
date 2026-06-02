from __future__ import annotations

from datetime import datetime
from io import BytesIO

from docx import Document


def build_text_docx(
    *,
    title: str,
    content: str,
) -> bytes:
    document = Document()

    document.add_heading(
        title,
        level=1,
    )

    document.add_paragraph(
        f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )

    document.add_paragraph("")

    clean_content = (
        content or ""
    ).strip()

    if not clean_content:
        clean_content = (
            "Sin contenido para exportar."
        )

    for paragraph in clean_content.splitlines():
        paragraph = paragraph.strip()

        if paragraph:
            document.add_paragraph(
                paragraph
            )

    buffer = BytesIO()

    document.save(buffer)

    buffer.seek(0)

    return buffer.read()
