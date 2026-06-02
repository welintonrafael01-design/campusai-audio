from __future__ import annotations

from io import BytesIO

from pptx import Presentation
from pptx.util import Inches, Pt


def _clean_lines(content: str) -> list[str]:
    return [
        line.strip()
        for line in (content or "").splitlines()
        if line.strip()
    ]


def build_summary_pptx(
    *,
    title: str,
    content: str,
) -> bytes:
    presentation = Presentation()

    title_slide_layout = presentation.slide_layouts[0]
    bullet_slide_layout = presentation.slide_layouts[1]

    slide = presentation.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = title or "StudyBook AI"
    slide.placeholders[1].text = "Presentación generada automáticamente"

    lines = _clean_lines(content)

    if not lines:
        lines = ["Sin contenido disponible para generar la presentación."]

    chunk_size = 5

    for index in range(0, len(lines), chunk_size):
        chunk = lines[index:index + chunk_size]

        slide = presentation.slides.add_slide(bullet_slide_layout)
        slide.shapes.title.text = (
            "Puntos clave"
            if index == 0
            else f"Puntos clave {index // chunk_size + 1}"
        )

        body = slide.placeholders[1].text_frame
        body.clear()

        for line_index, line in enumerate(chunk):
            paragraph = (
                body.paragraphs[0]
                if line_index == 0
                else body.add_paragraph()
            )
            paragraph.text = line[:180]
            paragraph.level = 0
            paragraph.font.size = Pt(22)

    buffer = BytesIO()
    presentation.save(buffer)
    buffer.seek(0)

    return buffer.read()
